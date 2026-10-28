#!/usr/bin/env python3
"""꺼내는 단어장 — 학교별 학생 페이지 빌드 도구

  python3 _build/build.py import <단어장.xlsx> <데이터이름>        엑셀 → _build/data/<데이터이름>.json (단어 그룹)
  python3 _build/build.py import-order <본문.txt> <데이터이름>     순서 텍스트 → _build/data/<데이터이름>.json (순서 그룹)
  python3 _build/build.py import-kw <핵심어.txt> <데이터이름>       핵심어 텍스트 → _build/data/<데이터이름>.json (핵심어 그룹)
  python3 _build/build.py <학교>                                   _build/schools/<학교>.json → 저장소 맨 위 <page>.html
  python3 _build/build.py demo                                     템플릿의 예시 세트로 선생님 확인용 index.html

세트 ID(단어: id_prefix-첫번호-끝번호, 순서: id_prefix-과키, 핵심어: id_prefix-강키)는 학생 진행 기록의 열쇠다. 학생에게 링크가
나간 뒤에는 id_prefix나 set_size(단어), 과 키(순서)·강 키(핵심어)를 바꾸지 않는다(바꾸면 그 그룹의 진행 기록이 처음부터 다시 시작됨).
뒤쪽 세트 크기만 바꿀 때는 set_size를 목록으로 써서 앞 세트를 그대로 둔다: [80, 160] = 첫 세트 80개, 그다음부터 160개씩.

순서 텍스트 형식 (_build/src/<데이터이름>.txt):
  # L5 | 5과 · From Trash to Music | 설명        ← 과(= 세트 하나). 키(세트 ID에 들어감) | 세트 제목 | 카드 설명(선택, 없으면 그룹 note나 '대화문 5 + 본문 4')
  ## Listen & Speak 1 ① | 대화문 | 환영 → 감사 → … ← 지문. 이름 | 종류(대화문/본문) | 흐름 요약(오답 때 보여 줌, 선택)
  Hey, Bora. Welcome to our rock band. | 첫 문장: …  ← 문장 | 왜 이 자리인지 한국어 단서(선택). 말한 사람 표시는 쓰지 않음
  ...

핵심어 텍스트 형식 (_build/src/<데이터이름>.txt):
  # 16 | 16강                                       ← 세트 하나(강). '|' 앞이 세트 ID에 들어가는 키
  ## 16강 01                                        ← 지문 이름
  = 한 줄 주제(한국어, **핵심어** 굵게) | 영어      ← 채점 뒤 보여 줌
  > 본문 문장                                       ← 차례대로(번호는 자동). 요약문은 넣지 않음
  + friction | 이유                                 ← 핵심 보기. '- 보기 | 이유'는 오답(지엽·함정)
  - pedal [pedalling] | 이유                        ← [ ] = 본문에서 찾을 꼴(보기와 꼴이 다를 때)
  + counteract | 이유 | 흐름                        ← 끝 칸: 동의어·흐름 = 본문에 없는 말(없으면 본문에 있는 말)
  * 메모                                            ← '* '로 시작하는 줄은 무시
"""
import sys, os, re, json, glob

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TEMPLATE = os.path.join(HERE, 'template.html')
DEFAULT_BATCH = 20  # template.html의 MAXB와 같은 값. 그룹 batch가 이 값이면 페이지 데이터에 따로 적지 않음

# 모든 학교가 함께 쓰는 제출 폼(구글폼 '단어 숙제 제출'). 학교 칸으로 학교를 구분한다.
FORM = {
    'url': 'https://docs.google.com/forms/d/e/1FAIpQLSfsvGXKVGASob6ylSlyCFKyEahrzGSqFNtwlaYEvA1AHOanww/viewform',
    'name': 'entry.2087311015', 'school': 'entry.331616395', 'set': 'entry.1245213658',
    'time': 'entry.1880995760', 'hard': 'entry.770504430',
}

HEADERS = {  # 머리글 이름 → 칸 (엑셀마다 열 순서가 달라도 머리글로 찾음)
    'n': {'번호', 'no', 'no.', '#'},
    'en': {'영어', 'english', 'word'},
    'ko': {'단어 뜻', '뜻', '외울 뜻', 'meaning'},
    'note': {'참고', '본문 풀이', '풀이', '비고'},
    'sec': {'지문', '출처', '구분', '지문 구분', 'passage'},
}
HANGUL = re.compile(r'[가-힯]')


def cell(v):
    return '' if v is None else str(v).strip()


# ---------- 엑셀 → 데이터 ----------
def import_xlsx(path, name):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['단어장'] if '단어장' in wb.sheetnames else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = next((i for i, r in enumerate(rows) if any(cell(c) == '영어' for c in r)), None)
    if hi is None:
        sys.exit('머리글 행(‘영어’ 칸)을 못 찾았어. 엑셀 머리글을 확인해 줘.')
    col = {}
    for j, c in enumerate(rows[hi]):
        h = cell(c).lower()
        for k, names in HEADERS.items():
            if h in names and k not in col: col[k] = j
    if 'en' not in col or 'ko' not in col:
        sys.exit(f'영어·뜻 칸을 못 찾았어. 머리글: {[cell(c) for c in rows[hi]]}')
    pre = [cell(c) for r in rows[:hi] for c in r if cell(c)]
    title = next((t for t in pre if '어학원' not in t), pre[0] if pre else '')  # 학원 이름 줄 다음의 제목 줄

    items, sec, run = [], None, 0
    for r in rows[hi + 1:]:
        get = lambda k: cell(r[col[k]]) if k in col and col[k] < len(r) else ''
        en, ko = get('en'), get('ko')
        filled = [cell(c) for c in r if cell(c)]
        if not en and not ko:
            if len(filled) == 1 and not filled[0].startswith('▶'):
                sec = filled[0]  # 지문 구분 행 (예: [교재 | Lesson 3 | 제목])
            continue
        if not en or not ko:
            print(f'  ! 영어나 뜻이 비어서 뺀 줄: {filled}')
            continue
        run += 1
        n = get('n')
        items.append({'n': int(float(n)) if re.fullmatch(r'\d+(\.0)?', n) else run, 'en': en, 'ko': ko,
                      'note': get('note'), 'sec': get('sec') or sec})

    nums = [it['n'] for it in items]
    gaps = [k for k in range(nums[0], nums[-1] + 1) if k not in set(nums)] if nums else []
    dup_n = sorted({k for k in nums if nums.count(k) > 1})
    low = [it['en'].lower() for it in items]
    dup_en = sorted({w for w in low if low.count(w) > 1})
    by_ko = {}
    for it in items: by_ko.setdefault(it['ko'], []).append(it['en'])
    twins = {k: v for k, v in by_ko.items() if len(v) > 1}
    no_eq = lambda en: re.sub(r'\s*\(\s*=[^)]*\)', '', en)  # (=다른 표현)은 한→영 단어 수에 안 셈 (template.html variants와 같음)
    long_en = [it['en'] for it in items if len([w for w in no_eq(it['en']).split() if not re.fullmatch(r'[A-Z~…+]', w) and not HANGUL.search(w)]) > 7]
    eq_en = [it['en'] for it in items if no_eq(it['en']) != it['en']]

    out = os.path.join(HERE, 'data', name + '.json')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, 'w', encoding='utf-8') as f:
        json.dump({'source': os.path.basename(path), 'title': title, 'items': items}, f, ensure_ascii=False, indent=1)
    secs = []
    for it in items:
        if not secs or secs[-1][0] != it['sec']: secs.append([it['sec'], 0])
        secs[-1][1] += 1
    print(f'{out}: 단어 {len(items)}개 (번호 {nums[0] if nums else "-"}~{nums[-1] if nums else "-"})')
    print(f'  제목: {title}')
    print('  지문: ' + (' / '.join(f'{(s or "구분 없음")[:40]} {c}' for s, c in secs) or '구분 없음'))
    if gaps or dup_n: print(f'  ! 번호 빠짐 {gaps[:10]} / 중복 {dup_n[:10]}')
    if dup_en: print(f'  ! 같은 표제어 {dup_en[:10]}')
    if twins: print(f'  · 뜻이 똑같은 단어(한→영에서 둘 다 정답 처리): {list(twins.values())[:10]}')
    if long_en: print(f'  · 8단어 이상이라 한→영 철자 문제에서 빠지는 표제어: {long_en[:10]}')
    if eq_en: print(f'  · (=…) 다른 표현도 한→영 정답으로 받는 표제어 {len(eq_en)}개: {eq_en[:6]}{" …" if len(eq_en) > 6 else ""}')


# ---------- 순서 텍스트 → 데이터 ----------
def import_order(path, name):
    lessons, cur_l, cur_p = [], None, None
    for raw in open(path, encoding='utf-8'):
        line = raw.strip()
        if not line: continue
        if line.startswith('## '):
            if cur_l is None: sys.exit(f'지문 앞에 과 줄(# 키 | 제목)이 없어: {line}')
            parts = [p.strip() for p in line[3:].split('|')]
            cur_p = {'name': parts[0], 'kind': parts[1] if len(parts) > 1 else ('본문' if parts[0].startswith('본문') else '대화문'),
                     'flow': parts[2] if len(parts) > 2 else '', 'lines': [], 'cues': []}
            cur_l['passages'].append(cur_p)
        elif line.startswith('# '):
            parts = [p.strip() for p in line[2:].split('|')]
            if len(parts) < 2 or not re.fullmatch(r'[A-Za-z0-9_-]+', parts[0]):
                sys.exit(f'과 줄은 "# 키 | 제목" 형식이어야 해 (키는 영문·숫자): {line}')
            cur_l = {'key': parts[0], 'title': parts[1], 'note': parts[2] if len(parts) > 2 else '', 'passages': []}; cur_p = None
            lessons.append(cur_l)
        else:
            if cur_p is None: sys.exit(f'문장 앞에 지문 줄(## 이름 | 종류)이 없어: {line}')
            sent, _, cue = line.partition(' | ')  # '문장 | 왜 이 자리인지 한국어 단서' (단서는 선택)
            cur_p['lines'].append(sent.strip()); cur_p['cues'].append(cue.strip())

    bad = []
    keys = [l['key'] for l in lessons]
    if len(set(keys)) != len(keys): bad.append(f'과 키 중복: {keys}')
    for l in lessons:
        names = [p['name'] for p in l['passages']]
        if len(set(names)) != len(names): bad.append(f"{l['key']} 지문 이름 중복: {names}")
        for p in l['passages']:
            if len(p['lines']) < 2: bad.append(f"{l['key']} {p['name']}: 문장이 {len(p['lines'])}개뿐")
            low = [s.lower() for s in p['lines']]
            if len(set(low)) != len(low): bad.append(f"{l['key']} {p['name']}: 똑같은 문장이 두 번 (순서를 맞출 수 없음)")
            for s in p['lines']:
                if re.match(r'^(?:[A-Z]|[A-Z][a-z]+|[가-힣]{1,3}):\s', s): bad.append(f"{l['key']} {p['name']}: 말한 사람 표시가 남아 있음 → {s[:40]}")
            n_cue = sum(1 for c in p['cues'] if c)
            if 0 < n_cue < len(p['lines']): bad.append(f"{l['key']} {p['name']}: 단서가 {len(p['lines'])}문장 중 {n_cue}개에만 있음")
            if n_cue and not p['flow']: bad.append(f"{l['key']} {p['name']}: 문장 단서는 있는데 흐름(## 줄 3번째 칸)이 없음")
    out = os.path.join(HERE, 'data', name + '.json')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, 'w', encoding='utf-8') as f:
        json.dump({'source': os.path.basename(path), 'lessons': lessons}, f, ensure_ascii=False, indent=1)
    for l in lessons:
        cued = sum(1 for p in l['passages'] if any(p['cues']))
        print(f"{out if l is lessons[0] else ''}\n  {l['key']} {l['title']}: " + ' / '.join(f"{p['name']} {len(p['lines'])}" for p in l['passages'])
              + f" → {len(l['passages'])}지문 {sum(len(p['lines']) for p in l['passages'])}문장 · 해설 있는 지문 {cued}")
    if bad: print('  ! ' + '\n  ! '.join(bad))


# ---------- 핵심어 텍스트 → 데이터 ----------
KW_OPT = re.compile(r'^([+-])\s+(.+?)(?:\s+\[([^\]]+)\])?\s+\|\s+(.+?)(?:\s+\|\s+(지문|동의어|흐름))?$')
KW_CATS = ('지문', '동의어', '흐름')


def kw_found(lines, form):
    pat = re.compile(r'(?<![A-Za-z-])' + re.escape(form) + r'(?![A-Za-z-])', re.I)
    return [i + 1 for i, s in enumerate(lines) if pat.search(s)]


def kw_inflect(w):
    out = [w, w + 's', w + 'es', w + 'd', w + 'ed', w + 'ing']
    if w.endswith('e'): out.append(w[:-1] + 'ing')
    if w.endswith('y'): out += [w[:-1] + 'ies', w[:-1] + 'ied']
    return out


def import_kw(path, name):
    sets, cur_s, cur_p = [], None, None
    for no, raw in enumerate(open(path, encoding='utf-8'), 1):
        line = raw.strip()
        if not line or line.startswith('*'): continue
        if line.startswith('## '):
            if cur_s is None: sys.exit(f'{no}줄: 지문 앞에 세트 줄(# 키 | 제목)이 없어: {line}')
            cur_p = {'name': line[3:].strip(), 'ko': '', 'en': '', 'lines': [], 'opts': [], 'forms': []}
            cur_s['passages'].append(cur_p)
        elif line.startswith('# '):
            parts = [p.strip() for p in line[2:].split('|')]
            if len(parts) < 2 or not re.fullmatch(r'[A-Za-z0-9_-]+', parts[0]):
                sys.exit(f'{no}줄: 세트 줄은 "# 키 | 제목" 형식이어야 해 (키는 영문·숫자): {line}')
            cur_s = {'key': parts[0], 'title': parts[1], 'passages': []}; cur_p = None
            sets.append(cur_s)
        else:
            if cur_p is None: sys.exit(f'{no}줄: 지문 줄(## 이름) 앞에 내용이 있어: {line}')
            if line.startswith('= '):
                ko, _, en = line[2:].partition(' | ')
                cur_p['ko'], cur_p['en'] = ko.strip(), en.strip()
            elif line.startswith('> '):
                cur_p['lines'].append(line[2:].strip())
            elif line[:1] in '+-':
                m = KW_OPT.match(line)
                if not m: sys.exit(f'{no}줄: 보기 줄은 "+ 보기 | 이유" 또는 "- 보기 | 이유 | 동의어" 형식이어야 해: {line}')
                sign, word, forms, why, cat = m.groups()
                cur_p['opts'].append([word.strip(), 1 if sign == '+' else 0, why.strip(), cat or '지문'])
                cur_p['forms'].append([f.strip() for f in forms.split(',')] if forms else [word.strip()])
            else:
                sys.exit(f'{no}줄: 무슨 줄인지 모르겠어 (=, >, +, -, * 중 하나로 시작해야 함): {line}')

    bad = []
    keys = [s['key'] for s in sets]
    if len(set(keys)) != len(keys): bad.append(f'세트 키 중복: {keys}')
    for s in sets:
        names = [p['name'] for p in s['passages']]
        if len(set(names)) != len(names): bad.append(f"{s['key']} 지문 이름 중복: {names}")
        for p in s['passages']:
            tag = f"{s['key']} {p['name']}"
            if not p['ko'] or not p['en']: bad.append(f'{tag}: 한 줄 주제(= 한국어 | 영어)가 없음')
            for t in (p['ko'], p['en']):
                if t.count('**') % 2: bad.append(f'{tag}: 주제 줄의 ** 짝이 안 맞음 → {t[:40]}')
            if len(p['lines']) < 2: bad.append(f"{tag}: 본문 문장이 {len(p['lines'])}개뿐")
            n, nk = len(p['opts']), sum(o[1] for o in p['opts'])
            if not 12 <= n <= 15: bad.append(f'{tag}: 보기가 {n}개 (12~15개로 정함)')
            if not nk or nk == n: bad.append(f'{tag}: 핵심 {nk}개 / 보기 {n}개 — 핵심과 오답이 둘 다 있어야 함')
            low = [o[0].lower() for o in p['opts']]
            if len(set(low)) != len(low): bad.append(f'{tag}: 같은 보기가 두 번')
            flows = sorted(o[1] for o in p['opts'] if o[3] == '흐름')
            if flows and flows != [0, 1]: bad.append(f'{tag}: 흐름 어휘는 정답 1 + 함정 1이어야 함 (지금 {flows})')
            for o, forms in zip(p['opts'], p['forms']):
                if not o[2]: bad.append(f'{tag}: {o[0]} 이유가 비어 있음')
                if o[3] == '지문':
                    if not any(kw_found(p['lines'], f) for f in forms):
                        bad.append(f"{tag}: '{o[0]}'가 본문에 없음 — 꼴이 다르면 [형태]로 적거나 '| 동의어'로 표시")
                else:
                    hit = [(f, kw_found(p['lines'], f)) for f in kw_inflect(o[0])]
                    hit = [(f, h) for f, h in hit if h]
                    if hit: bad.append(f"{tag}: '{o[0]}'({o[3]})가 본문에 있음 → {hit[0][0]} {hit[0][1]}번")
            p.pop('forms')
    out = os.path.join(HERE, 'data', name + '.json')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, 'w', encoding='utf-8') as f:
        json.dump({'source': os.path.basename(path), 'sets': sets}, f, ensure_ascii=False, indent=1)
    print(out)
    for s in sets:
        print(f"  {s['key']} {s['title']}: " + ' / '.join(f"{p['name'].split()[-1]} {sum(o[1] for o in p['opts'])}/{len(p['opts'])}" for p in s['passages'])
              + f" → {len(s['passages'])}지문 (정답 수/보기 수)")
    if bad: print('  ! ' + '\n  ! '.join(bad))


def kw_sets(g):
    data = json.load(open(os.path.join(HERE, 'data', g['data'] + '.json'), encoding='utf-8'))
    want = g.get('units')
    sets = []
    for u in data['sets']:
        if want and u['key'] not in want: continue
        ps = u['passages']
        sizes = [len(p['opts']) for p in ps]
        names = [p['name'] for p in ps]
        span = (f"{names[0]}~{names[-1].split()[-1]}번" if len(names) > 1 and all(len(x.split()) == 2 for x in names)
                and len({x.split()[0] for x in names}) == 1 else f'{len(ps)}지문')
        st = {'id': f"{g['id_prefix']}-{u['key']}", 'kind': 'kw', 'group': g['group'], 'title': u['title'],
              'note': g.get('note') or f"{span} · 지문마다 보기 {min(sizes)}~{max(sizes)}개", 'passages': ps}
        sets.append(st)
        print(f"{g['group'][:24]} | {st['title']} | {st['note']} | {len(ps)}지문 · 정답 " + '·'.join(str(sum(o[1] for o in p['opts'])) for p in ps))
    return sets


def order_sets(g):
    data = json.load(open(os.path.join(HERE, 'data', g['data'] + '.json'), encoding='utf-8'))
    want = g.get('lessons')
    sets = []
    for l in data['lessons']:
        if want and l['key'] not in want: continue
        kinds = []
        for p in l['passages']:
            if not kinds or kinds[-1][0] != p['kind']: kinds.append([p['kind'], 0])
            kinds[-1][1] += 1
        st = {'id': f"{g['id_prefix']}-{l['key']}", 'kind': 'order', 'group': g['group'], 'title': l['title'],
              'note': l.get('note') or g.get('note') or ' + '.join(f'{k} {c}' for k, c in kinds), 'passages': l['passages']}
        sets.append(st)
        print(f"{g['group'][:24]} | {st['title']} | {st['note']} | {len(l['passages'])}지문 {sum(len(p['lines']) for p in l['passages'])}문장")
    return sets


# ---------- 데이터 → 학교 페이지 ----------
def sec_name(raw, cont, default):
    if not raw: return default or '단어'
    if re.fullmatch(r'\d+강 \d+', raw.strip()): return raw.strip() + '번'  # 지문 칸 '16강 01' → 카드에 '16강 01번 44'
    parts = [p.strip() for p in raw.strip('[]').split('|')]
    if len(parts) < 3: return raw.strip('[]')
    lesson, name = parts[1], parts[2]
    if name.startswith('대화문'): return f'{lesson} 대화문'
    return f'{lesson} 본문 (이어서)' if cont else f'{lesson} 본문 · {name}'


def build(school):
    cfg_path = os.path.join(HERE, 'schools', school + '.json')
    cfg = json.load(open(cfg_path, encoding='utf-8'))
    sets = []
    for g in [g for g in cfg['groups'] if g.get('active', True)]:
        if g.get('type') == 'order':
            sets += order_sets(g); continue
        if g.get('type') == 'kw':
            sets += kw_sets(g); continue
        items = json.load(open(os.path.join(HERE, 'data', g['data'] + '.json'), encoding='utf-8'))['items']
        # set_size: 숫자면 모두 그 크기, 목록이면 차례대로 쓰고 마지막 값을 반복 ([80, 160] → 80, 160, 160, …)
        sizes = g['set_size'] if isinstance(g['set_size'], list) else [g['set_size']]
        bounds, i = [], 0
        while i < len(items):
            size = int(sizes[min(len(bounds), len(sizes) - 1)])
            bounds.append((i, min(i + size, len(items)))); i += size
        if len(bounds) > 1 and bounds[-1][1] - bounds[-1][0] < size // 2:  # 자투리(그 세트 크기의 절반 미만)는 앞 세트에 합침
            bounds[-2] = (bounds[-2][0], bounds[-1][1]); bounds.pop()
        seen = set()
        for k, (lo, hi) in enumerate(bounds, 1):
            chunk = items[lo:hi]; n0, n1 = chunk[0]['n'], chunk[-1]['n']
            # "rounds": true면 단어장 종이의 회차를 제목에 붙임 (올림포스 3회차 · 187~279번). 제목만 바뀌고 세트 ID는 그대로
            label = f"{g['title']} {k}회차" if g.get('rounds') else g['title']
            st = {'id': f"{g['id_prefix']}-{n0}-{n1}", 'group': g['group'], 'title': f"{label} · {n0}~{n1}번", 'sections': []}
            if int(g.get('batch', DEFAULT_BATCH)) != DEFAULT_BATCH: st['batch'] = int(g['batch'])
            for it in chunk:
                if not st['sections'] or st['sections'][-1]['raw'] != it['sec']:
                    st['sections'].append({'raw': it['sec'], 'name': sec_name(it['sec'], it['sec'] in seen, g.get('note')), 'items': []})
                st['sections'][-1]['items'].append([it['en'], it['ko'], it['note']])
            seen.update(s['raw'] for s in st['sections'])
            st['note'] = g.get('note') or ' + '.join(  # 지문 하나('16강 05번 42')는 줄이 안 갈라지게 붙임
                f"{s['name'].split(' · ')[0].replace(' (이어서)', '')} {len(s['items'])}".replace(' ', ' ') for s in st['sections'])
            for s in st['sections']: s.pop('raw')
            sets.append(st)
            print(f"{g['group'][:24]} | {st['title']} | {st['note']} | {hi - lo}개")

    page_cfg = {'paste': False, 'school': cfg['school'], 'fields': cfg['fields'], 'form': FORM}
    kinds = {s.get('kind', 'vocab') for s in sets}  # 순서 세트만 / 핵심어 세트만 있는 페이지는 탭 제목도 그 이름
    name = '본문 순서 맞추기' if kinds == {'order'} else '핵심어 고르기' if kinds == {'kw'} else '꺼내는 단어장'
    write_page(cfg['page'], sets, page_cfg, f"{name} · {cfg['short']}")


def write_page(page, sets, page_cfg, title):
    """template.html을 완성된 HTML 문서로 감싸서 저장소 맨 위에 씀. sets가 None이면 템플릿의 예시 세트를 그대로 둠"""
    tpl = open(TEMPLATE, encoding='utf-8').read()
    if sets is not None:
        data = json.dumps(sets, ensure_ascii=False, indent=1)
        assert '</script' not in data
        tpl = re.sub(r'(<script id="sets-data" type="application/json">)[\s\S]*?(</script>)',
                     lambda m: m.group(1) + '\n' + data + '\n' + m.group(2), tpl, count=1)
    tpl = tpl.replace('<script id="sets-data"', f'<script id="page-config" type="application/json">{json.dumps(page_cfg, ensure_ascii=False)}</script>\n<script id="sets-data"', 1)
    tpl = tpl.replace('<title>꺼내는 단어장</title>', f'<title>{title}</title>', 1)
    tpl = re.sub(r'<meta name="claude:viewport-fit-sides"[^>]*>\n?', '', tpl)
    tpl = tpl.replace(':root{\n  --paper', ':root{\n  color-scheme:light;\n  --paper', 1)
    cut = tpl.index('</style>') + len('</style>')
    head, body = tpl[:cut], tpl[cut:].strip()
    out = ('<!doctype html>\n<html lang="ko">\n<head>\n<meta charset="utf-8">\n'
           '<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">\n'
           '<meta name="robots" content="noindex,nofollow">\n'
           '<style>:root{padding-top:env(safe-area-inset-top,0px);padding-bottom:env(safe-area-inset-bottom,0px)}body{margin:0}img{max-width:100%}[hidden]{display:none!important}</style>\n'
           + head + '\n</head>\n<body>\n' + body + '\n</body>\n</html>\n')
    path = os.path.join(ROOT, page)
    open(path, 'w', encoding='utf-8').write(out)
    print(f"→ {page} ({len(out.encode()) // 1024} KB) · 링크 https://kiarakimmy71617.github.io/vocab/{page}")


def build_demo():
    """선생님 확인용 데모(index.html): 템플릿의 예시 세트(단어 + 순서) 그대로, 단어 직접 넣기 켜짐"""
    write_page('index.html', None, {'paste': True, 'school': '', 'fields': []}, '꺼내는 단어장 · 데모')


if __name__ == '__main__':
    if len(sys.argv) == 4 and sys.argv[1] == 'import':
        import_xlsx(sys.argv[2], sys.argv[3])
    elif len(sys.argv) == 4 and sys.argv[1] == 'import-order':
        import_order(sys.argv[2], sys.argv[3])
    elif len(sys.argv) == 4 and sys.argv[1] == 'import-kw':
        import_kw(sys.argv[2], sys.argv[3])
    elif len(sys.argv) == 2 and sys.argv[1] == 'demo':
        build_demo()
    elif len(sys.argv) == 2:
        build(sys.argv[1])
    else:
        sys.exit(__doc__)
